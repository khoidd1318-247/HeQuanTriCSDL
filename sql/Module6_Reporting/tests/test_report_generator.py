"""
Unit tests for Module 6 - Reporting (Excel & PDF generation)
Compatible with both pytest and python -m unittest
"""
import sys
from pathlib import Path
import unittest
import openpyxl

MODULE_ROOT = Path(__file__).resolve().parent.parent
if str(MODULE_ROOT) not in sys.path:
    sys.path.insert(0, str(MODULE_ROOT))

from src.report_generator import (
    get_default_traffic_data,
    get_default_model_comparison_data,
    export_excel_report,
    export_pdf_report,
)


class TestReportGenerator(unittest.TestCase):
    def setUp(self):
        self.test_dir = MODULE_ROOT / "tests" / "test_outputs"
        self.test_dir.mkdir(parents=True, exist_ok=True)
        self.excel_path = self.test_dir / "test_report.xlsx"
        self.pdf_path = self.test_dir / "test_report.pdf"

    def tearDown(self):
        # Cleanup temporary files
        if self.excel_path.exists():
            self.excel_path.unlink()
        if self.pdf_path.exists():
            self.pdf_path.unlink()
        if self.test_dir.exists():
            try:
                self.test_dir.rmdir()
            except OSError:
                pass

    def test_get_default_data(self):
        meta, df_ts, df_hotspots = get_default_traffic_data()
        self.assertEqual(meta["RecordingID"], "REC_IND_01")
        self.assertEqual(len(df_ts), 10)
        self.assertEqual(len(df_hotspots), 5)

        df_models = get_default_model_comparison_data()
        self.assertEqual(len(df_models), 2)
        self.assertIn("Mô hình", df_models.columns)

    def test_export_excel_report(self):
        res_path = export_excel_report(self.excel_path)
        self.assertTrue(Path(res_path).exists())
        self.assertGreater(Path(res_path).stat().st_size, 0)

        # Inspect workbook sheets
        wb = openpyxl.load_workbook(res_path)
        expected_sheets = ["Tong_Quan", "Chuoi_Thoi_Gian", "Diem_Nong_Un_Tac", "Doi_Sanh_Mo_Hinh"]
        for s in expected_sheets:
            self.assertIn(s, wb.sheetnames)

        ws1 = wb["Tong_Quan"]
        self.assertIn("BÁO CÁO PHÂN TÍCH GIAO THÔNG", str(ws1.cell(row=1, column=1).value))

    def test_export_pdf_report(self):
        res_path = export_pdf_report(self.pdf_path)
        self.assertTrue(Path(res_path).exists())
        self.assertGreater(Path(res_path).stat().st_size, 1000)


if __name__ == "__main__":
    unittest.main()
